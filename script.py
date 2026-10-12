from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
from datetime import datetime
from pathlib import Path
import re
import requests
from bs4 import BeautifulSoup

# ============================================================
# EINSTELLUNGEN
# ============================================================

EXCEL_DATEI = Path("Systemanforderung.xlsx")
TABELLENBLATT = "Tabelle1"


DATEV_URL = (
    "https://www.datev.de/web/de/berufsgruppenuebergreifend/"
    "service-und-support/rund-um-unsere-programme/"
    "technische-informationen/systemplattform/"
    "systemvoraussetzungen-fuer-datev-software"
)

SAGE_URL = "https://sage-news.de/systemcheck"

SAGE_VERSION_URL = (
    "https://www.1st-vision.de/portfolio/sage-100/"
    "liveupdate-changelog/"
)

RAMICRO_URL = (
    "https://helpdesk.ra-micro.de/help/de-de/"
    "639-systemvoraussetzungen/"
    "385-ra-micro-systemvoraussetzungen"
)


# ============================================================
# HILFSFUNKTIONEN
# ============================================================

def sage_text_bereinigen(text):
    """Entfernt überflüssige Leerzeichen."""
    return " ".join(str(text).split())


def sage_webseite_lesen(url):
    """Lädt eine Webseite herunter und gibt ihren Text zurück."""

    response = requests.get(
        url,
        headers={
            "User-Agent": "Mozilla/5.0"
        },
        timeout=60
    )

    response.raise_for_status()

    soup = BeautifulSoup(response.text, "html.parser")

    for element in soup(["script", "style", "noscript"]):
        element.decompose()

    return soup.get_text("\n", strip=True)


def sage_wert_suchen(text, muster):
    """Sucht mithilfe eines regulären Ausdrucks nach einem Wert."""

    treffer = re.search(
        muster,
        text,
        re.IGNORECASE
    )

    if treffer:
        return sage_text_bereinigen(treffer.group(1))

    return None



# ============================================================
# DATEV
# ============================================================

def datev_auslesen(browser):

    print("\n========================================")
    print(" DATEV")
    print("========================================")

    page = browser.new_page()

    try:

        print("DATEV-Hauptseite öffnen...")

        page.goto(
            DATEV_URL,
            wait_until="domcontentloaded",
            timeout=60000
        )

        # Cookie-Banner
        try:

            cookie = page.get_by_role(
                "button",
                name="Nur erforderliche Cookies"
            )

            if cookie.count() > 0:
                cookie.click()

        except Exception:
            pass


        # Hardware-Dokument suchen
        print("Hardware-Dokument suchen...")

        hardware_link = page.get_by_role(
            "link",
            name=re.compile(
                "Hardware-Voraussetzungen",
                re.IGNORECASE
            )
        ).first

        hardware_link.wait_for(
            state="visible",
            timeout=30000
        )

        hardware_link.click()

        page.wait_for_load_state(
            "domcontentloaded"
        )


        # RAM-Zelle dient als Anker
        ram_element = page.get_by_text(
            "Arbeitsspeicher / RAM",
            exact=True
        ).first

        ram_element.wait_for(
            state="visible",
            timeout=60000
        )

        # Tabelle finden
        tabelle = ram_element.locator(
            "xpath=ancestor::table"
        )

        zeilen = tabelle.locator("tr")

        daten = []

        for i in range(zeilen.count()):

            zellen = zeilen.nth(i).locator(
                "th, td"
            )

            werte = []

            for j in range(zellen.count()):

                werte.append(
                    sage_text_bereinigen(
                        zellen.nth(j).inner_text()
                    )
                )

            if werte:
                daten.append(werte)


        # DATEV-geeignet-Spalte finden
        datev_spalte = None

        for zeile in daten:

            for index, wert in enumerate(zeile):

                if "DATEV-geeignet" in wert:

                    datev_spalte = index
                    break

            if datev_spalte is not None:
                break


        if datev_spalte is None:

            raise Exception(
                "DATEV-geeignet-Spalte nicht gefunden."
            )


        cpu = None
        ram = None
        storage = None


        for zeile in daten:

            if not zeile:
                continue

            name = zeile[0].lower()

            if len(zeile) <= datev_spalte:
                continue


            if (
                "prozessor" in name
                or "cpu" in name
            ):
                cpu = zeile[datev_spalte]


            elif (
                "arbeitsspeicher" in name
                or "ram" in name
            ):
                ram = zeile[datev_spalte]


            elif (
                "festplatte" in name
                or "speicherplatz" in name
                or "datenträger" in name
            ):
                datenträger = zeile[datev_spalte]


        print(f"CPU:     {cpu}")
        print(f"RAM:     {ram}")
        print(f"DATENTRÄGER: {datenträger}")


        return {
            "Programm": "DATEV",
            "Erweitert": "DATEV - RDS",
            "CPU": cpu,
            "RAM": ram,
            "DATENTRÄGER": datenträger
        }


    finally:

        page.close()


# ============================================================
# SAGE 100 RDS
# ============================================================

# ============================================================
# SAGE - SOFTWAREVERSION PRÜFEN
# ============================================================

def sage_version_pruefen():

    print("\n[1] Sage-Version überprüfen...")

    try:
        text = sage_webseite_lesen(SAGE_VERSION_URL)

        # Sucht Versionsnummern wie 9.0.11.5
        versionen = re.findall(
            r"\b\d+\.\d+\.\d+(?:\.\d+)?\b",
            text
        )

        if not versionen:
            print("    WARNUNG: Keine Version gefunden.")
            return None

        # Höchste gefundene Versionsnummer
        version = max(
            versionen,
            key=lambda v: tuple(map(int, v.split(".")))
        )

        print(f"    Gefundene Version: {version}")

        return version

    except Exception as fehler:
        print(f"    Versionsprüfung fehlgeschlagen: {fehler}")
        return None


# ============================================================
# SAGE - HARDWAREANFORDERUNGEN AUSLESEN
# ============================================================
import re


def bereinigen(text):
    """
    Entfernt überflüssige Leerzeichen und Zeilenumbrüche.
    """

    if text is None:
        return ""

    return " ".join(str(text).split())


def wert_aus_text(text, pattern):
    """
    Sucht mit einem regulären Ausdruck nach einem Wert.
    """

    treffer = re.search(
        pattern,
        text,
        re.IGNORECASE
    )

    if treffer:
        return bereinigen(treffer.group(1))

    return None

def sage_auslesen(browser):

    print("\n========================================")
    print(" SAGE 100 - RDS")
    print("========================================")

    aktuelle_version = sage_version_pruefen()

    print("\n[2] Systemanforderungen herunterladen...")

    text = sage_webseite_lesen(SAGE_URL)

    print(f"    {len(text)} Zeichen ausgelesen.")

    # --------------------------------------------------------
    # Terminalserver-Bereich suchen
    # --------------------------------------------------------

    print("\n[3] Terminalserver-Bereich suchen...")

    start = re.search(
        r"Terminal[\s-]*Server",
        text,
        re.IGNORECASE
    )

    if not start:
        raise Exception(
            "Terminalserver-Bereich wurde nicht gefunden."
        )

    terminal_text = text[start.start():]

    # Nachfolgende Bereiche ausschließen
    ende = re.search(
        r"Application[\s-]*Server",
        terminal_text,
        re.IGNORECASE
    )

    if ende:
        terminal_text = terminal_text[:ende.start()]

    print("    Terminalserver-Bereich gefunden.")

    # --------------------------------------------------------
    # CPU
    # --------------------------------------------------------

    print("\n[4] CPU-Anforderungen auslesen...")

    cpu = sage_wert_suchen(
        terminal_text,
        r"CPU\s*Min:\s*(.*?)\s+Empf:"
    )

    # --------------------------------------------------------
    # RAM
    # --------------------------------------------------------

    print("[5] RAM-Anforderungen auslesen...")

    ram_pro_sitzung = sage_wert_suchen(
        terminal_text,
        r"RAM\s+pro\s+Sitzung\s+Standardbetrieb:"
        r"\s*([0-9]+(?:[.,][0-9]+)?\s*GB)"
    )

    ram_basis = sage_wert_suchen(
        terminal_text,
        r"Formel:\s*"
        r"([0-9]+(?:[.,][0-9]+)?\s*GB)\s*Basis"
    )

    ram = None

    if ram_basis and ram_pro_sitzung:
        ram = (
            f"{ram_basis} Basis + "
            f"{ram_pro_sitzung} je Benutzer"
        )

    # --------------------------------------------------------
    # STORAGE
    # --------------------------------------------------------

    print("[6] Storage-Anforderungen auslesen...")

    storage = sage_wert_suchen(
        terminal_text,
        r"Freier\s+Speicher\s+Min:\s*(.*?)\s+Empf:"
    )

    # --------------------------------------------------------
    # ERGEBNIS
    # --------------------------------------------------------

    print("\n========================================")
    print(" SAGE 100 - ERGEBNIS")
    print("========================================")

    print(f"Softwareversion: {aktuelle_version or 'Unbekannt'}")
    print(f"CPU:             {cpu or 'Nicht gefunden'}")
    print(f"RAM:             {ram or 'Nicht gefunden'}")
    print(f"Storage:         {storage or 'Nicht gefunden'}")

    # --------------------------------------------------------
    # PLAUSIBILITÄTSPRÜFUNG
    # --------------------------------------------------------

    if not all([cpu, ram, storage]):
        print(
            "\nWARNUNG: Nicht alle Hardwareanforderungen "
            "konnten ausgelesen werden."
        )

        print(
            "Die vorhandenen Excel-Werte dürfen für "
            "fehlende Angaben nicht überschrieben werden."
        )

 

    print("\n========================================")
    print(" SAGE 100 - RDS")
    print("========================================")

    page = browser.new_page()

    try:
        print("[1] Sage-Webseite öffnen...")

        print(f"Verwendete Sage-URL: {SAGE_URL}")

        page.goto(
            SAGE_URL,
            wait_until="domcontentloaded",
            timeout=60000
        )

        print(f"    Geladene URL: {page.url}")

        # Warten, bis der Seiteninhalt vorhanden ist
        page.locator("body").wait_for(
            state="visible",
            timeout=30000
        )

        print("[2] Seiteninhalt auslesen...")

        text = page.locator("body").inner_text()

        if not text:
            raise Exception(
                "Die Sage-Webseite hat keinen Text geliefert."
            )

        print(f"    {len(text)} Zeichen geladen.")

        # ----------------------------------------------------
        # Terminalserver-Bereich suchen
        # ----------------------------------------------------

        print("[3] Terminalserver-Bereich suchen...")

        start = text.lower().find("terminalserver")

        if start == -1:
            raise Exception(
                "Der Begriff 'Terminalserver' wurde "
                "auf der Sage-Webseite nicht gefunden."
            )

        # Der Terminalserver-Bereich endet beim
        # nächsten Application-Server-Abschnitt.
        ende = text.lower().find(
            "application-server",
            start
        )

        if ende == -1:
            terminal_text = text[start:]
        else:
            terminal_text = text[start:ende]

        print("    Terminalserver-Bereich gefunden.")

        # Debug-Ausgabe
        print("\n--- Gefundener Sage Terminalserver-Text ---")
        print(terminal_text)
        print("-------------------------------------------")

        # ----------------------------------------------------
        # CPU
        # ----------------------------------------------------

        print("\n[4] CPU suchen...")

        cpu_match = re.search(
            r"CPU\s*Min:\s*(.+?)(?:\r?\n|\s+Empf:)",
            terminal_text,
            re.IGNORECASE
        )

        if cpu_match:
            cpu = sage_text_bereinigen(cpu_match.group(1))
        else:
            cpu = None

        # ----------------------------------------------------
        # RAM pro Sitzung
        # ----------------------------------------------------

        print("[5] RAM suchen...")

        ram_sitzung_match = re.search(
            r"RAM\s+pro\s+Sitzung\s+Standardbetrieb:\s*"
            r"([0-9]+(?:[.,][0-9]+)?\s*GB)",
            terminal_text,
            re.IGNORECASE
        )

        if ram_sitzung_match:
            ram_sitzung = sage_text_bereinigen(
                ram_sitzung_match.group(1)
            )
        else:
            ram_sitzung = None

        # Basis-RAM des Servers
        ram_basis_match = re.search(
            r"Formel:\s*"
            r"([0-9]+(?:[.,][0-9]+)?\s*GB)\s*Basis",
            terminal_text,
            re.IGNORECASE
        )

        if ram_basis_match:
            ram_basis = sage_text_bereinigen(
                ram_basis_match.group(1)
            )
        else:
            ram_basis = None

        if ram_sitzung and ram_basis:
            ram = (
                f"{ram_basis} Basis + "
                f"{ram_sitzung} je Benutzer"
            )

        elif ram_sitzung:
            ram = f"{ram_sitzung} je Benutzer"

        else:
            ram = None

        # ----------------------------------------------------
        # STORAGE
        # ----------------------------------------------------

        print("[6] Storage suchen...")

        storage_match = re.search(
            r"Freier\s+Speicher\s+Min:\s*"
            r"(.+?)(?:\r?\n|\s+Empf:)",
            terminal_text,
            re.IGNORECASE
        )

        if storage_match:
            storage = sage_text_bereinigen(
                storage_match.group(1)
            )
        else:
            storage = None

        # ----------------------------------------------------
        # Ergebnis
        # ----------------------------------------------------

        print("\n========================================")
        print(" SAGE ERGEBNIS")
        print("========================================")

        print(f"CPU:     {cpu}")
        print(f"RAM:     {ram}")
        print(f"Storage: {storage}")

        if cpu is None:
            print("WARNUNG: CPU konnte nicht gefunden werden.")

        if ram is None:
            print("WARNUNG: RAM konnte nicht gefunden werden.")

        if storage is None:
            print("WARNUNG: Storage konnte nicht gefunden werden.")

        return {
            "Programm": "Sage",
            "Erweitert": "Sage 100 - RDS",
            "CPU": cpu,
            "RAM": ram,
            "Storage": storage
        }

    finally:
        page.close()





# ============================================================
# RA-MICRO RDS
# ============================================================

def ramicro_auslesen(browser):

    print("\n========================================")
    print(" RA-MICRO - RDS")
    print("========================================")

    page = browser.new_page()

    try:

        print("RA-MICRO öffnen...")

        page.goto(
            RAMICRO_URL,
            wait_until="domcontentloaded",
            timeout=60000
        )

        page.wait_for_timeout(2000)

        text = page.locator("body").inner_text()

        text = sage_text_bereinigen(text)


        # ----------------------------------------------------
        # Terminal-Server-Bereich
        # ----------------------------------------------------

        start = text.lower().find(
            "terminal-server"
        )

        if start == -1:

            start = text.lower().find(
                "terminalserver"
            )


        if start == -1:

            raise Exception(
                "Terminal-Server-Bereich bei "
                "RA-MICRO nicht gefunden."
            )


        # Nächster Bereich
        ende = text.lower().find(
            "hinweis für profilcontainer",
            start
        )


        if ende == -1:

            ende = text.lower().find(
                "pc-desktop",
                start
            )


        if ende == -1:

            terminal_text = text[start:]

        else:

            terminal_text = text[start:ende]


        # ----------------------------------------------------
        # CPU
        # ----------------------------------------------------

        cpu = sage_wert_suchen(
            terminal_text,
            r"CPU:\s*(.*?)(?=\s+RAM:)"
        )


        # ----------------------------------------------------
        # RAM
        # ----------------------------------------------------

        ram = sage_wert_suchen(
            terminal_text,
            r"RAM:\s*(.*?)(?=\s+bei Nutzung|\s+virtueller)"
        )


        # ----------------------------------------------------
        # STORAGE
        # ----------------------------------------------------

        datenträger = wert_aus_text(
            terminal_text,
            r"Speicher:\s*(.*?)(?=\s+Netzwerk:)"
        )


        print(f"CPU:     {cpu}")
        print(f"RAM:     {ram}")
        print(f"DATENTRÄGER: {datenträger}")


        return {
            "Programm": "RA-Micro",
            "Erweitert": "RA-Micro - RDS",
            "CPU": cpu,
            "RAM": ram,
            "DATENTRÄGER": datenträger
        }


    finally:

        page.close()


# ============================================================
# EXCEL-ZEILE FINDEN
# ============================================================

def excel_zeile_finden(
    ws,
    programm,
    erweitert
):

    for row in range(
        1,
        ws.max_row + 1
    ):

        excel_programm = bereinigen(
            ws.cell(
                row=row,
                column=1
            ).value
        )

        excel_erweitert = bereinigen(
            ws.cell(
                row=row,
                column=2
            ).value
        )


        if (
            excel_programm.lower()
            == programm.lower()
            and
            excel_erweitert.lower()
            == erweitert.lower()
        ):

            return row


    return None


# ============================================================
# EINZELNEN WERT PRÜFEN
# ============================================================

def wert_pruefen(
    ws,
    row,
    column,
    bezeichnung,
    neuer_wert
):

    alter_wert = bereinigen(
        ws.cell(
            row=row,
            column=column
        ).value
    )

    neuer_wert = bereinigen(
        neuer_wert
    )


    print()
    print(bezeichnung)

    print(
        f"  ALT: {alter_wert}"
    )

    print(
        f"  NEU: {neuer_wert}"
    )


    # --------------------------------------------------------
    # Sicherheitsfunktion:
    # Leere Web-Werte niemals nach Excel schreiben
    # --------------------------------------------------------

    if not neuer_wert:

        print(
            "  -> Kein Web-Wert gefunden."
        )

        print(
            "  -> Excel bleibt unverändert."
        )

        return False


    if alter_wert == neuer_wert:

        print(
            "  -> unverändert"
        )

        return False


    print(
        "  -> ÄNDERUNG ERKANNT"
    )


    ws.cell(
        row=row,
        column=column
    ).value = neuer_wert


    return True


# ============================================================
# HERSTELLER IN EXCEL AKTUALISIEREN
# ============================================================

def hersteller_aktualisieren(
    ws,
    daten
):

    programm = daten["Programm"]
    erweitert = daten["Erweitert"]


    print()
    print("-" * 60)

    print(
        f"Excel prüfen: {programm} / {erweitert}"
    )

    print("-" * 60)


    row = excel_zeile_finden(
        ws,
        programm,
        erweitert
    )


    if row is None:

        print(
            "Excel-Zeile wurde nicht gefunden!"
        )

        return 0


    print(
        f"Gefunden in Excel-Zeile {row}."
    )


    aenderungen = 0


    # C = CPU
    if wert_pruefen(
        ws,
        row,
        3,
        "CPU",
        daten["CPU"]
    ):

        aenderungen += 1


    # D = RAM
    if wert_pruefen(
        ws,
        row,
        4,
        "RAM",
        daten["RAM"]
    ):

        aenderungen += 1


    # E = Storage
    if wert_pruefen(
        ws,
        row,
        5,
        "Datenträger",
        daten["Storage"]
    ):

        aenderungen += 1


    # --------------------------------------------------------
    # Datum nur bei tatsächlicher Änderung
    # --------------------------------------------------------

    if aenderungen > 0:

        datum = datetime.now().strftime(
            "%d.%m.%Y"
        )

        # F = letzte Aktualisierung
        ws.cell(
            row=row,
            column=6
        ).value = datum


        print()
        print(
            f"{aenderungen} Änderung(en)."
        )

        print(
            f"Aktualisierungsdatum: {datum}"
        )


    else:

        print()
        print(
            "Keine Änderungen gefunden."
        )


    return aenderungen


# ============================================================
# HAUPTPROGRAMM
# ============================================================

def main():

    print()
    print("####################################################")
    print("#                                                  #")
    print("#       SYSTEMANFORDERUNGEN UPDATE                 #")
    print("#                                                  #")
    print("#       DATEV | SAGE 100 | RA-MICRO               #")
    print("#                                                  #")
    print("####################################################")


    # --------------------------------------------------------
    # Excel prüfen
    # --------------------------------------------------------

    if not EXCEL_DATEI.exists():

        print()
        print(
            "FEHLER:"
        )

        print(
            f"{EXCEL_DATEI} wurde nicht gefunden."
        )

        return


    # --------------------------------------------------------
    # Webseiten auslesen
    # --------------------------------------------------------

    alle_daten = []


    with sync_playwright() as p:

        browser = p.chromium.launch(
            headless=True
        )


        # ----------------------------------------------------
        # DATEV
        # ----------------------------------------------------

        try:

            daten = datev_auslesen(
                browser
            )

            alle_daten.append(
                daten
            )

        except Exception as e:

            print()
            print(
                "DATEV konnte nicht "
                "ausgelesen werden:"
            )

            print(e)


        # ----------------------------------------------------
        # SAGE
        # ----------------------------------------------------

        try:

            daten = sage_auslesen(
                browser
            )

            alle_daten.append(
                daten
            )

        except Exception as e:

            print()
            print(
                "Sage konnte nicht "
                "ausgelesen werden:"
            )

            print(e)


        # ----------------------------------------------------
        # RA-MICRO
        # ----------------------------------------------------

        try:

            daten = ramicro_auslesen(
                browser
            )

            alle_daten.append(
                daten
            )

        except Exception as e:

            print()
            print(
                "RA-MICRO konnte nicht "
                "ausgelesen werden:"
            )

            print(e)


        browser.close()


    # --------------------------------------------------------
    # Keine Daten?
    # --------------------------------------------------------

    if not alle_daten:

        print()
        print(
            "Keine Herstellerdaten konnten "
            "ausgelesen werden."
        )

        return


    # --------------------------------------------------------
    # Excel öffnen
    # --------------------------------------------------------

    wb = load_workbook(
        EXCEL_DATEI
    )


    if TABELLENBLATT not in wb.sheetnames:

        print(
            f"Tabellenblatt "
            f"'{TABELLENBLATT}' "
            f"nicht gefunden."
        )

        wb.close()

        return


    ws = wb[TABELLENBLATT]


    # --------------------------------------------------------
    # Excel aktualisieren
    # --------------------------------------------------------

    gesamt_aenderungen = 0


    for daten in alle_daten:

        gesamt_aenderungen += (
            hersteller_aktualisieren(
                ws,
                daten
            )
        )


    # --------------------------------------------------------
    # Speichern
    # --------------------------------------------------------

    if gesamt_aenderungen > 0:

        try:

            wb.save(
                EXCEL_DATEI
            )

        except PermissionError:

            print()
            print(
                "FEHLER:"
            )

            print(
                "Systemanforderung.xlsx "
                "ist vermutlich noch in "
                "Excel geöffnet."
            )

            print(
                "Bitte Excel schließen und "
                "das Script erneut starten."
            )

            wb.close()

            return


    wb.close()


    # --------------------------------------------------------
    # Ergebnis
    # --------------------------------------------------------

    print()
    print("=" * 60)

    if gesamt_aenderungen > 0:

        print(
            f" {gesamt_aenderungen} "
            f"ÄNDERUNG(EN) GEFUNDEN"
        )

        print()
        print(
            " Systemanforderung.xlsx "
            "wurde aktualisiert."
        )

    else:

        print(
            " KEINE ÄNDERUNGEN GEFUNDEN"
        )

        print()
        print(
            " Excel wurde nicht verändert."
        )

    print("=" * 60)


# ============================================================
# START
# ============================================================

if __name__ == "__main__":
    main()